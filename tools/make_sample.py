#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
make_sample.py — 生成示例简历（DOCX + PDF 双格式，内容一致）

用于验证转换工具的两条解析链路：
  - DOCX：带标题样式（Heading），走样式标记识别
  - PDF ：纯文本层，走正则模式识别
"""
import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
SAMPLES_DIR = os.path.abspath(os.path.join(HERE, "..", "samples"))
os.makedirs(SAMPLES_DIR, exist_ok=True)

# ---------------------------------------------------------------- 简历内容（虚构人物）
NAME = "陈雨桐"
HEADLINE = "前端开发工程师"
CONTACT = "电话：138-0013-8000 ｜ 邮箱：chenyutong.dev@example.com ｜ GitHub：github.com/cytong ｜ 博客：https://chenyutong.dev"

SUMMARY = ("3 年前端开发经验，专注 Vue/React 生态与前端工程化。主导过多个从 0 到 1 的中后台项目，"
           "擅长性能优化与组件体系建设，有跨端（Web/小程序）与 Node.js 全栈实践。")

EDUCATION = [
    ("2016.09 - 2020.06", "同济大学", "本科 · 计算机科学与技术",
     "GPA 3.6/4.0；主修课程：数据结构、操作系统、计算机网络、数据库原理；校级优秀学生奖学金 ×2"),
]

EXPERIENCE = [
    ("2022.03 - 至今", "字节跳动", "高级前端开发工程师",
     ["主导抖音电商商家后台性能优化，首屏加载时间从 3.2s 降至 1.1s（LCP）",
      "搭建基于 Arco Design 的二次封装组件库，覆盖 12 条业务线、复用率 78%",
      "推动团队前端工程化升级：Vite 迁移、Monorepo 改造、CI 流水线标准化"]),
    ("2020.07 - 2022.02", "美团", "前端开发工程师",
     ["负责美团外卖商家端 H5 活动页与营销工具开发，日均 PV 300 万+",
      "设计活动页可视化搭建 Schema，运营自助配置活动占比从 15% 提升至 60%"]),
]

PROJECTS = [
    ("简历海 resume-sea", "统一的简历发布平台，支持 PDF/Word 简历一键转换发布",
     ["设计 RSUF 统一简历数据格式与模板适配层，解析准确率约 85%",
      "基于 Jekyll 主题生态实现多模板渲染，发布即得独立 URL"]),
    ("个人博客系统", "基于 Next.js + MDX 的技术博客，月活 2 万+",
     ["实现静态生成 + ISR，Lighthouse 性能分 98",
      "自研评论与订阅系统，接入 GitHub OAuth"]),
]

SKILLS = [
    "精通：Vue3、TypeScript、JavaScript、HTML5/CSS3",
    "熟练：React、Node.js、Vite、Webpack、Sass",
    "熟悉：小程序开发、Vitest 单元测试、Docker",
    "了解：Python、Nginx、基础 CI/CD",
]

CERTS = [
    ("AWS Certified Cloud Practitioner", "Amazon", "2023.05"),
    ("同济大学优秀毕业生", "同济大学", "2020.06"),
]

LANGUAGES = "中文（母语）、英语（CET-6，可流畅阅读技术文档）"


# ---------------------------------------------------------------- DOCX 生成
def make_docx(path):
    import docx
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = docx.Document()

    # 姓名
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(NAME)
    run.font.size = Pt(22)
    run.font.bold = True

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(HEADLINE)
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(0x55, 0x5E, 0x6C)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(CONTACT)
    run.font.size = Pt(9)

    def heading(text):
        h = doc.add_heading(text, level=1)
        for r in h.runs:
            r.font.size = Pt(14)

    heading("个人简介")
    doc.add_paragraph(SUMMARY)

    heading("教育经历")
    for time_range, school, degree, detail in EDUCATION:
        doc.add_paragraph(f"{time_range}  {school}  {degree}")
        doc.add_paragraph(detail)

    heading("工作经历")
    for time_range, company, role, details in EXPERIENCE:
        doc.add_paragraph(f"{time_range}  {company}  {role}")
        for d in details:
            doc.add_paragraph(f"· {d}")

    heading("项目经历")
    for name, tagline, details in PROJECTS:
        doc.add_paragraph(name)
        doc.add_paragraph(tagline)
        for d in details:
            doc.add_paragraph(f"· {d}")

    heading("专业技能")
    for line in SKILLS:
        doc.add_paragraph(line)

    heading("证书与荣誉")
    for name, org, date in CERTS:
        doc.add_paragraph(f"{name}  {org}  {date}")

    heading("语言")
    doc.add_paragraph(LANGUAGES)

    doc.save(path)
    print("生成 DOCX:", path)


# ---------------------------------------------------------------- PDF 生成
def find_cjk_font():
    candidates = [
        r"C:\Windows\Fonts\Deng.ttf",
        r"C:\Windows\Fonts\Dengb.ttf",
        r"C:\Windows\Fonts\ARIALUNI.ttf",
        r"C:\Windows\Fonts\simsun.ttc",
        r"C:\Windows\Fonts\msyh.ttc",
    ]
    for c in candidates:
        if os.path.isfile(c):
            return c
    return None


def make_pdf(path):
    from fpdf import FPDF

    font = find_cjk_font()
    if not font:
        print("[跳过] 未找到中文字体，跳过 PDF 生成", file=sys.stderr)
        return

    pdf = FPDF()
    pdf.add_page()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_font("cjk", "", font)
    pdf.add_font("cjk", "B", font)

    def center_line(text, font_style, size, height, color=None):
        """手动居中（fpdf2 的 multi_cell align=C 配合部分中文字体存在宽度计算 bug）"""
        pdf.set_font("cjk", font_style, size)
        if color:
            pdf.set_text_color(*color)
        w = pdf.get_string_width(text)
        pdf.set_x(max(pdf.l_margin, (pdf.w - w) / 2))
        pdf.multi_cell(0, height, text)

    def title(text):
        center_line(text, "B", 20, 10, (30, 30, 30))

    def sub(text):
        center_line(text, "", 10, 6, (90, 94, 108))

    def heading(text):
        pdf.ln(4)
        pdf.set_font("cjk", "B", 13)
        pdf.set_text_color(30, 30, 30)
        pdf.multi_cell(0, 8, text)

    def body(text, size=10, indent=0):
        pdf.set_font("cjk", "", size)
        pdf.set_text_color(40, 40, 40)
        pdf.set_x(pdf.l_margin + indent)
        pdf.multi_cell(0, 6, text)

    title(NAME)
    sub(HEADLINE)
    sub(CONTACT)

    heading("个人简介")
    body(SUMMARY)

    heading("教育经历")
    for time_range, school, degree, detail in EDUCATION:
        body(f"{time_range}  {school}  {degree}")
        body(detail, size=9)

    heading("工作经历")
    for time_range, company, role, details in EXPERIENCE:
        body(f"{time_range}  {company}  {role}")
        for d in details:
            body(f"· {d}", size=9, indent=4)

    heading("项目经历")
    for name, tagline, details in PROJECTS:
        body(name)
        body(tagline, size=9)
        for d in details:
            body(f"· {d}", size=9, indent=4)

    heading("专业技能")
    for line in SKILLS:
        body(line, size=9)

    heading("证书与荣誉")
    for name, org, date in CERTS:
        body(f"{name}  {org}  {date}", size=9)

    heading("语言")
    body(LANGUAGES, size=9)

    pdf.output(path)
    print("生成 PDF:", path)


if __name__ == "__main__":
    make_docx(os.path.join(SAMPLES_DIR, "sample-resume.docx"))
    make_pdf(os.path.join(SAMPLES_DIR, "sample-resume.pdf"))
